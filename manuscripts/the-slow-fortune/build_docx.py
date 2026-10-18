"""Assemble book/part-*.md (SPEC markup) into a styled Word workbook."""
import glob
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SRC = sys.argv[1]
OUT = sys.argv[2]
RUNNING = sys.argv[3] if len(sys.argv) > 3 else "Beating Your Neighbor or Beating the Clock"
DOC_TITLE = sys.argv[4] if len(sys.argv) > 4 else RUNNING

INK = RGBColor(0x1F, 0x2A, 0x44)
ACCENT = RGBColor(0x8A, 0x5A, 0x1E)
BODY_FONT = "Georgia"
HEAD_FONT = "Calibri"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Inches(1))
sec.top_margin = sec.bottom_margin = Inches(0.9)

st = doc.styles
st["Normal"].font.name = BODY_FONT
st["Normal"].font.size = Pt(10.5)
st["Normal"].paragraph_format.space_after = Pt(6)
st["Normal"].paragraph_format.line_spacing = 1.15
for name, size, color in (("Heading 1", 22, INK), ("Heading 2", 16, INK), ("Heading 3", 12.5, ACCENT)):
    s = st[name]
    s.font.name = HEAD_FONT
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = color
    rfonts = s.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), HEAD_FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    s.paragraph_format.space_before = Pt(14 if name != "Heading 1" else 0)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.keep_with_next = True

# Footer: running title + page number
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run(RUNNING + "  •  ")
r.font.size = Pt(8.5)
r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
r2 = fp.add_run()
r2.font.size = Pt(8.5)
for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
    if kind:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
    else:
        el = OxmlElement("w:instrText")
        el.set(qn("xml:space"), "preserve")
        el.text = text
    r2._r.append(el)

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*)")


def add_runs(par, text, size=None, color=None, italic=False, bold=False):
    for piece in INLINE.split(text):
        if not piece:
            continue
        b, i = bold, italic
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            piece, b = piece[2:-2], True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            piece, i = piece[1:-1], True
        run = par.add_run(piece)
        run.bold, run.italic = b or None, i or None
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hex_fill)
    tcPr.append(sh)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def write_lines(n, container=None):
    target = container or doc
    for _ in range(n):
        p = target.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(10)
        pPr = p._p.get_or_add_pPr()
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:leader"), "underscore")
        tab.set(qn("w:pos"), str(int(6.5 * 1440)))
        tabs.append(tab)
        pPr.append(tabs)
        r = p.add_run("\t")
        r.font.color.rgb = RGBColor(0xA6, 0xA6, 0xA6)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def write_table(rows):
    header, body = rows[0], rows[1:]
    ncol = len(header)
    t = doc.add_table(rows=1, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        c.paragraphs[0].text = ""
        add_runs(c.paragraphs[0], h, size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(c, "1F2A44")
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    trPr.append(th)
    for row in body:
        cells = t.add_row().cells
        filled = any(x for x in row)
        for j in range(ncol):
            val = row[j] if j < len(row) else ""
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, val, size=9.5)
        if not filled or sum(1 for x in row if x) <= 1:
            t.rows[-1].height = Inches(0.42)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def write_box(title, lines):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    shade(cell, "F4EEE4")
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    for k, v in (("w:val", "single"), ("w:sz", "24"), ("w:color", "8A5A1E")):
        left.set(qn(k), v)
    borders.append(left)
    tcPr.append(borders)
    p = cell.paragraphs[0]
    add_runs(p, title, size=12.5, bold=True, color=ACCENT)
    p.paragraph_format.keep_with_next = True
    para = []

    def flush():
        if para:
            q = cell.add_paragraph()
            add_runs(q, " ".join(para), size=10.5)
            para.clear()

    for ln in lines:
        if not ln.strip():
            flush()
        elif re.match(r"^\s*[-*] ", ln):
            flush()
            q = cell.add_paragraph()
            q.paragraph_format.left_indent = Inches(0.25)
            add_runs(q, "•  " + re.sub(r"^\s*[-*] ", "", ln), size=10.5)
        elif re.match(r"^#{2,4} ", ln):
            flush()
            q = cell.add_paragraph()
            add_runs(q, re.sub(r"^#+ ", "", ln), size=11, bold=True)
        else:
            para.append(ln.strip())
    flush()
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def title_page(lines):
    for _ in range(7):
        doc.add_paragraph()
    for idx, ln in enumerate(l for l in lines if l.strip()):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        text = ln.strip().strip("#").strip().strip("*")
        if idx == 0:
            add_runs(p, text, size=44, bold=True, color=INK)
        elif idx == 1:
            add_runs(p, text, size=15, italic=True, color=ACCENT)
        else:
            if idx == 2:
                doc.add_paragraph()
                doc.add_paragraph()
            add_runs(p, text, size=11, color=RGBColor(0x55, 0x55, 0x55))


files = sorted(glob.glob(f"{SRC}/part-*.md"))
first = True
for path in files:
    lines = open(path, encoding="utf-8").read().replace("\r\n", "\n").split("\n")
    in_chapter = False
    i, para = 0, []

    def flush_para():
        if para:
            p = doc.add_paragraph()
            add_runs(p, " ".join(s.strip() for s in para))
            para.clear()

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s in ("# FRONTMATTER", "# BACKMATTER"):
            flush_para()
            if not first:
                page_break()
            first = False
            i += 1
            continue
        if s == "[[TITLE]]":
            flush_para()
            j = i + 1
            block = []
            while j < len(lines) and lines[j].strip() != "[[PAGEBREAK]]":
                block.append(lines[j])
                j += 1
            title_page(block)
            i = j
            continue
        if s == "[[PAGEBREAK]]":
            flush_para()
            page_break()
            i += 1
            continue
        m = re.match(r"^\[\[LINES (\d+)\]\]$", s)
        if m:
            flush_para()
            write_lines(int(m.group(1)))
            i += 1
            continue
        if re.match(r"^# CHAPTER \d+", s):
            flush_para()
            if not first:
                page_break()
            first = False
            in_chapter = True
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(90)
            add_runs(p, s[2:].strip(), size=13, bold=True, color=ACCENT)
            i += 1
            continue
        if s.startswith("## TITLE:"):
            flush_para()
            h = doc.add_heading(s[len("## TITLE:"):].strip(), level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        if s.startswith(":::box"):
            flush_para()
            title = s[len(":::box"):].strip()
            j = i + 1
            body = []
            while j < len(lines) and lines[j].strip() != ":::":
                body.append(lines[j])
                j += 1
            write_box(title, body)
            i = j + 1
            continue
        if s.startswith("|"):
            flush_para()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = split_row(lines[i])
                is_sep = any(cells) and all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c)
                if not is_sep:
                    rows.append(cells)
                i += 1
            if rows:
                write_table(rows)
            continue
        if s.startswith("#"):
            flush_para()
            level = len(s) - len(s.lstrip("#"))
            text = s.lstrip("#").strip()
            if level == 1:
                doc.add_heading(text, level=1)
            elif level == 2:
                doc.add_heading(text, level=2 if in_chapter else 1)
            else:
                doc.add_heading(text, level=3)
            i += 1
            continue
        if s.startswith("> "):
            flush_para()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(10)
            add_runs(p, s[2:], size=13, italic=True, color=ACCENT)
            i += 1
            continue
        if re.match(r"^[-*] ", s):
            flush_para()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, "•  " + s[2:])
            i += 1
            continue
        m = re.match(r"^(\d+)\. (.*)", s)
        if m:
            flush_para()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, f"{m.group(1)}.  {m.group(2)}")
            i += 1
            continue
        if not s:
            flush_para()
            i += 1
            continue
        para.append(s)
        i += 1
    flush_para()

doc.core_properties.title = DOC_TITLE
doc.core_properties.author = ""
doc.save(OUT)
print("saved", OUT)
