#!/usr/bin/env python3
"""Core Technology Overview v15 = v14 (every paragraph unchanged) + Section 56, The Self-Knowledge Layer (SKL), and a
subtitle and closing note naming the addition. Additions of October 22, 2026 are highlighted bright green (v14's
October 8 additions stay yellow). Usage: build_overview_v15.py <v14.docx> <v15.docx>"""
import copy, os, sys
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.text.paragraph import Paragraph
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import skl_text as T

src, out = sys.argv[1:3]
d = Document(src); P = d.paragraphs
ORIG = [copy.deepcopy(p._p) for p in P]
G = WD_COLOR_INDEX.BRIGHT_GREEN
find = lambda s: next(p for p in P if p.text.startswith(s))
h1, body = find("55.  Hybrid Transport"), P[P.index(find("55.  Hybrid Transport")) + 1]
h2 = find("1.1  Precise Definition")
note = find("Additions of October 8, 2026")

def clone(tpl, text, bold=None):
    p = copy.deepcopy(tpl._p)
    para = Paragraph(p, tpl._parent)
    for r in para.runs[1:]: r._r.getparent().remove(r._r)
    r = para.runs[0]; r.text = text; r.font.highlight_color = G
    if bold is not None: r.bold = bold
    return p

els = [clone(h1, f"56.  {T.TITLE}")] + [clone(body, t) for t in T.INTRO]
for head, paras in T.SUBSECTIONS:
    els.append(clone(h2, head)); els += [clone(body, t) for t in paras]
anchor = note._p.getprevious()                     # the last paragraph of §55
for el in els:
    anchor.addnext(el); anchor = el
n2 = clone(note, f"Additions of {T.DATE}, highlighted in green: Section 56, The Self-Knowledge Layer (SKL), built in prototype v9.64 "
                 "(qb_selfknow.py) and the applications' shared core (qbselfknow.py); Composite Feature Registry v70 features 383–391; "
                 "provisional FIG. 57 and claims 358–362. No existing text was changed.")
note._p.addnext(n2)
sub = P[1]
r = sub.add_run(" · v15 addendum, " + T.DATE + ": Section 56 The Self-Knowledge Layer · Registry v70: 391 features · 362 claims")
r.italic = True; r.font.highlight_color = G
if sub.runs[0].font.size: r.font.size = sub.runs[0].font.size

# additive check: every v14 paragraph still present, in order, unchanged (the subtitle only gains a run)
from lxml import etree
now = [etree.tostring(copy.deepcopy(p._p)) for p in d.paragraphs]
old = [etree.tostring(p) for p in ORIG]
it = iter(now)
missing = [i for i, o in enumerate(old) if i != 1 and o not in it]
assert not missing, missing[:5]
d.save(out)
print("added paragraphs:", len(d.paragraphs) - len(ORIG))
