"""Standalone VAT feature registry (from registry v67 data)."""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
S = sys.argv[1]; sys.path.insert(0, f"{S}/reg"); import v67_data as D
NUM = {int(k): v for k, v in json.load(open(f"{S}/vatset/numerals.json")).items()}
FIGMAP = {56: 1, 57: 2}
wb = Workbook()
HF = Font(name="Arial", bold=True, color="FFFFFF", size=10); HFILL = PatternFill("solid", start_color="1F3864")
BF = Font(name="Arial", size=10); AL = Alignment(wrap_text=True, vertical="top")
th = Side(style="thin", color="A6A6A6"); BB = Border(left=th, right=th, top=th, bottom=th)
def sheet(ws, heads, rows):
    for i, (h, w) in enumerate(heads, 1):
        c = ws.cell(1, i, h); c.font, c.fill, c.alignment, c.border = HF, HFILL, Alignment(wrap_text=True, vertical="center"), BB
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, row in enumerate(rows, 2):
        for i, v in enumerate(row, 1):
            c = ws.cell(r, i, v); c.font, c.alignment, c.border = BF, AL, BB
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{len(rows) + 1}"
# ---- Features
ws = wb.active; ws.title = "Features"
rows = []
for k, (name, dom, figs, surface, role, func, rule, wexpr, status, module) in enumerate(D.NEW, 1):
    here = ", ".join(str(FIGMAP[f]) for f in figs if f in FIGMAP)
    other = ", ".join(str(f) for f in figs if f not in FIGMAP)
    for a, b in NUM.items(): func = func.replace(str(a), str(b))
    func = func.replace("(5620–5626)", "(120–126)")
    rows.append([f"VAT-{k:02d}", name, dom, here, ", ".join(map(str, figs)) + (f" (also cites QueryBook FIG. {other})" if other else ""),
                 role, func, rule, wexpr, status, module, 363 + k])
sheet(ws, [("ID", 8), ("Feature", 30), ("Domain", 22), ("FIG. (this application)", 11), ("FIG. (QueryBook application)", 18), ("Role", 14),
           ("Function", 60), ("Rule", 44), ("Worked expression", 26), ("Status", 11), ("Prototype module", 20), ("QueryBook Registry v67 No.", 11)], rows)
assert "56" not in " ".join(r[6] for r in rows if "(56" in r[6])
# ---- Algorithms
wa = wb.create_sheet("Algorithms")
arows = []
for k, (aid, name, feats, figs, inp, proc, outp, det, mod) in enumerate(D.ALGOS, 1):
    arows.append([f"VAT-ALG-{k}", aid, name, feats, ", ".join(str(FIGMAP[f]) for f in figs if f in FIGMAP), inp, proc, outp, det, mod])
sheet(wa, [("ID", 11), ("QueryBook ID", 10), ("Algorithm", 26), ("Features", 32), ("FIG.", 6), ("Inputs", 24), ("Procedure", 80), ("Output", 24),
           ("Deterministic", 12), ("Module", 20)], arows)
# ---- Reference numerals
EL = {100: "AI agent (unmodified)", 105: "Egress-lock boundary", 110: "VAT sidecar", 112: "Protocol adapters (MCP, A2A, HTTP)", 114: "Claim extraction",
      116: "Intent binding", 118: "Frame codec and signing", 120: "Encrypted connection (QUIC streams or TLS over TCP)", 122: "Control lane", 124: "Normal lane",
      126: "Bulk lane", 130: "Peer VAT sidecar", 140: "Destination (agent or tool)", 150: "UFCS verification service", 160: "Ledger and session certificates",
      170: "Control plane", 200: "Outbound agent message", 205: "Sentence classification", 206: "Non-fact labelled, never verified", 210: "Triple extraction and semantic fingerprint",
      220: "Fingerprint lookup in UFCS store", 225: "FQL contradiction query", 231: "VERIFIED outcome", 232: "CONTRADICTED outcome", 233: "UNKNOWN outcome",
      240: "Signed VERDICT frame and delivery policy", 245: "Action vs intent; result vs tool record", 250: "Ledger append (hash chain, Merkle checkpoints)",
      270: "Agent Session Certificate", 280: "Agent reputation (evidence pair)", 290: "Have/want exchange", 295: "Privacy mode (fingerprint-only)"}
inv = {v: k for k, v in NUM.items()}
assert set(EL) == set(inv), set(EL) ^ set(inv)
sheet(wb.create_sheet("Reference Numerals"), [("Numeral", 9), ("Element", 46), ("FIG.", 6), ("QueryBook numeral", 12)],
      [[n, EL[n], 1 if n < 200 else 2, inv[n]] for n in sorted(EL)])
# ---- Claims map
CL = [(1, "Independent (method)", "", "VAT-01, VAT-04, VAT-06, VAT-07, VAT-08, VAT-09, VAT-15, VAT-16"), (2, "Dependent", 1, "VAT-09"),
      (3, "Dependent", 1, "VAT-11, VAT-12"), (4, "Dependent", 1, "VAT-13"), (5, "Dependent", 1, "VAT-14"), (6, "Dependent", 1, "VAT-06"),
      (7, "Dependent", 1, "VAT-01, VAT-02, VAT-03"), (8, "Dependent", 1, "VAT-10"), (9, "Dependent", 1, "VAT-17, VAT-18"), (10, "Dependent", 1, "VAT-15"),
      (11, "Independent (computer-readable medium)", "1–10", "All")]
sheet(wb.create_sheet("Claims Map"), [("Claim", 7), ("Type", 30), ("Depends on", 11), ("QueryBook claim", 11), ("Features", 60)],
      [[c, t, d, 354 + c, f] for c, t, d, f in CL])
# ---- Summary
sm = wb.create_sheet("Summary", 0)
sm["A1"] = "Verified Agent Transport — Feature Registry"; sm["A1"].font = Font(name="Arial", bold=True, size=14)
sm["A2"] = "The QueryBook AI Agent Wrapper Service. Extracted from QueryBook Composite Feature Registry v67 (rows 364–381); numbering follows the standalone application."
sm["A2"].font = Font(name="Arial", italic=True, size=10, color="595959")
R = [("Features", "=COUNTA(Features!A2:A200)"), ("Features specified, not yet built", '=COUNTIF(Features!J2:J200,"Specified")'),
     ("Features with a worked expression", '=SUMPRODUCT(--(Features!I2:I200<>""))'), ("Features shown in FIG. 1", '=COUNTIF(Features!D2:D200,"*1*")'),
     ("Features shown in FIG. 2", '=COUNTIF(Features!D2:D200,"*2*")'), ("Algorithms", "=COUNTA(Algorithms!A2:A100)"),
     ("Reference numerals", "=COUNTA('Reference Numerals'!A2:A100)"), ("Claims", "=COUNTA('Claims Map'!A2:A100)"),
     ("Independent claims", "=COUNTIF('Claims Map'!B2:B100,\"Independent*\")")]
for c, v in (("A4", "Measure"), ("B4", "Value")):
    sm[c] = v; sm[c].font, sm[c].fill = HF, HFILL
for i, (a, f) in enumerate(R, 5):
    sm.cell(i, 1, a).font = BF; sm.cell(i, 2, f).font = BF
r0 = 5 + len(R) + 1
sm.cell(r0, 1, "Features by domain").font = Font(name="Arial", bold=True)
for i, d in enumerate(sorted({r[2] for r in rows}), r0 + 1):
    sm.cell(i, 1, d).font = BF; sm.cell(i, 2, f"=COUNTIF(Features!C2:C200,A{i})").font = BF
sm.column_dimensions["A"].width = 44; sm.column_dimensions["B"].width = 10
wb.save(sys.argv[2]); print(len(rows), "features;", len(arows), "algorithms")
