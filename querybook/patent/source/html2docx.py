#!/usr/bin/env python3
"""Convert the provisional-patent print HTML into a Word document.
Usage: html2docx.py <in.html> <svg_png_dir> <out.docx> <highlight:1|0>
Elements carrying class "add"/"addfig" (or inside one) get yellow highlight when highlight=1.
"""
import re, sys
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

src, pngdir, out, HL = sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4] == "1"
soup = BeautifulSoup(open(src, encoding="utf-8").read(), "lxml")
wrap = soup.find(class_="wrap")

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Inches(1))

FONT = "Times New Roman"
def set_font(style, size=None, bold=None, color=RGBColor(0, 0, 0)):
    style.font.name = FONT
    style.element.get_or_add_rPr()
    rfonts = style.element.rPr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); style.element.rPr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    if size: style.font.size = Pt(size)
    if bold is not None: style.font.bold = bold
    if color is not None: style.font.color.rgb = color

set_font(doc.styles["Normal"], 12)
doc.styles["Normal"].paragraph_format.space_after = Pt(6)
for name, size in (("Title", 15), ("Heading 1", 14), ("Heading 2", 12.5), ("Heading 3", 12)):
    st = doc.styles[name]; set_font(st, size, True)
    st.font.italic = False
    st.paragraph_format.space_before = Pt(14 if name != "Title" else 6)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True
# Title style has a bottom border in the default template; remove it
tp = doc.styles["Title"].element.pPr
if tp is not None:
    for b in tp.findall(qn("w:pBdr")): tp.remove(b)

# footer: page number; header: draft legend
def add_field(par, instr):
    r = par.add_run(); f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin"); r._r.append(f1)
    r = par.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r._r.append(it)
    r = par.add_run(); f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end"); r._r.append(f2)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_field(fp, "PAGE")
hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
hr = hp.add_run("DRAFT — NOT FILED — QueryBook Provisional Patent Application — rev. October 8, 2026"
                + (" — additions highlighted" if HL else ""))
hr.font.size = Pt(8); hr.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

def classes(t): return t.get("class", []) if isinstance(t, Tag) else []
def is_added(t):
    while isinstance(t, Tag):
        c = classes(t)
        if "add" in c or "addfig" in c: return True
        t = t.parent
    return False

BLOCK = {"p", "div", "h1", "h2", "h3", "h4", "table", "ul", "ol", "li"}

def shade_par(par, fill, left_color=None):
    pPr = par._p.get_or_add_pPr()
    if left_color:
        bdr = OxmlElement("w:pBdr"); l = OxmlElement("w:left")
        for k, v in (("w:val", "single"), ("w:sz", "24"), ("w:space", "6"), ("w:color", left_color)): l.set(qn(k), v)
        bdr.append(l); _ins(pPr, bdr)
    shd = OxmlElement("w:shd")
    for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", fill)): shd.set(qn(k), v)
    _ins(pPr, shd)

LATER = {qn("w:" + n) for n in ("tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct",
         "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
         "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap",
         "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange")}
def _ins(pPr, el):
    for i, ch in enumerate(pPr):
        if ch.tag in LATER:
            pPr.insert(i, el); return
    pPr.append(el)

def emit_inline(par, node, fmt):
    if isinstance(node, NavigableString):
        if node.__class__.__name__ in ("Comment", "Doctype"): return
        text = re.sub(r"\s+", " ", str(node))
        if not text: return
        r = par.add_run(text)
        r.bold = fmt.get("b") or None
        r.italic = fmt.get("i") or None
        if fmt.get("size"): r.font.size = Pt(fmt["size"])
        if fmt.get("color"): r.font.color.rgb = fmt["color"]
        if fmt.get("caps"): r.font.all_caps = True
        if HL and fmt.get("hl"): r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return
    if not isinstance(node, Tag): return
    if node.name in ("style", "script", "title", "svg"): return
    f = dict(fmt)
    c = classes(node)
    if node.name in ("b", "strong") or "n" in c or "cn" in c: f["b"] = True
    if node.name in ("i", "em"): f["i"] = True
    if "fix" in c: f.update(b=True, caps=True, size=8, color=RGBColor(0xA9, 0x79, 0x1C))
    if "add" in c or "addfig" in c: f["hl"] = True
    st = node.get("style", "")
    if "color:#b25d00" in st.replace(" ", ""): f["color"] = RGBColor(0xB2, 0x5D, 0x00)
    if node.name == "br":
        par.add_run().add_break(); return
    for ch in node.children: emit_inline(par, ch, f)
    if "cn" in c or "n" in c:
        pass

def tidy(par):
    runs = [r for r in par.runs if r.text]
    if runs:
        runs[0].text = runs[0].text.lstrip()
        runs[-1].text = runs[-1].text.rstrip()
    return par

def new_par(style=None):
    return doc.add_paragraph(style=style) if style else doc.add_paragraph()

def has_block(t): return any(isinstance(c, Tag) and c.name in BLOCK for c in t.children)

def para_from(t, style=None, align=None, fmt=None, size=None):
    if not t.get_text(strip=True): return None
    p = new_par(style)
    if align is not None: p.alignment = align
    f = dict(fmt or {}); f["hl"] = is_added(t)
    if size: f["size"] = size
    emit_inline(p, t, f)
    return tidy(p)

fig_i = 0
def figure(t):
    global fig_i
    fig_i += 1
    divs = [c for c in t.children if isinstance(c, Tag) and c.name == "div"]
    label, caption = divs[0], divs[-1]
    p = para_from(label, fmt={"b": True}); p.paragraph_format.keep_with_next = True
    ip = new_par(); ip.alignment = WD_ALIGN_PARAGRAPH.CENTER; ip.paragraph_format.keep_with_next = True
    ip.add_run().add_picture(f"{pngdir}/f{fig_i:02d}.png", width=Inches(6.4))
    para_from(caption, align=WD_ALIGN_PARAGRAPH.CENTER, size=10)

def table(t):
    rows = t.find_all("tr")
    ncol = max(len(r.find_all(["td", "th"])) for r in rows)
    tb = doc.add_table(rows=0, cols=ncol); tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(0.6), Inches(2.65), Inches(0.6), Inches(2.65)] if ncol == 4 else [Inches(6.5 / ncol)] * ncol
    for r in rows:
        cells = r.find_all(["td", "th"])
        row = tb.add_row()
        for i, cell in enumerate(cells):
            dc = row.cells[i]; dc.width = widths[i]
            p = dc.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            f = {"size": 9.5, "hl": is_added(cell), "b": cell.name == "th"}
            emit_inline(p, cell, f); tidy(p)
        if r.find("th"):  # repeat header row
            trPr = row._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    tb.autofit = False
    for gc, w in zip(tb._tbl.tblGrid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(int(w.inches * 1440)))
    doc.add_paragraph()

def block(t):
    if isinstance(t, NavigableString):
        if str(t).strip() and t.__class__.__name__ == "NavigableString":
            p = new_par(); emit_inline(p, t, {"hl": is_added(t.parent)}); tidy(p)
        return
    if not isinstance(t, Tag) or t.name in ("style", "script", "title"): return
    c = classes(t); st = (t.get("style") or "").replace(" ", "")
    if t.name == "h1":
        p = para_from(t, "Title", WD_ALIGN_PARAGRAPH.CENTER)
        for r in p.runs: r.font.all_caps = True
    elif t.name == "h2": para_from(t, "Heading 1")
    elif t.name == "h3": para_from(t, "Heading 2")
    elif t.name == "h4": para_from(t, "Heading 3")
    elif t.name == "table": table(t)
    elif t.name in ("ul", "ol"):
        for li in t.find_all("li", recursive=False): para_from(li, "List Bullet" if t.name == "ul" else "List Number")
    elif t.name == "p":
        p = para_from(t, align=WD_ALIGN_PARAGRAPH.JUSTIFY if ({"pp", "claim"} & set(c)) else None)
        if p is not None and "font-size:12px" in st:
            for r in p.runs: r.font.size = Pt(10)
    elif t.name == "div":
        if "page-break-before:always" in st and not t.get_text(strip=True):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE); return
        if t.find("svg") and not t.find("div", class_="foot"):
            figure(t); return
        if "draft" in c:
            p = para_from(t, align=WD_ALIGN_PARAGRAPH.CENTER, fmt={"b": True, "caps": True, "color": RGBColor(0xB2, 0x5D, 0x00)}, size=9.5)
            return
        if "covermeta" in c:
            p = para_from(t, align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5)
            if p is not None and "uppercase" in st:
                for r in p.runs: r.font.all_caps = True
            return
        if "foot" in c and not has_block(t):
            p = para_from(t, size=9.5)
            if p is not None:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                if "border-left" in st: shade_par(p, "FBF7F0", "B25D00")
            return
        if has_block(t):
            for ch in t.children: block(ch)
        else:
            para_from(t)

for ch in wrap.children: block(ch)
doc.core_properties.title = "QueryBook Provisional Patent Application — rev. October 8, 2026"
doc.core_properties.author = "QueryBook Core Technology Group"
z = doc.settings.element.find(qn("w:zoom"))
if z is not None and z.get(qn("w:percent")) is None: z.set(qn("w:percent"), "100")
doc.save(out)
print("saved", out, "figures", fig_i)
