#!/usr/bin/env python3
"""Build the Excel registry (numerals, figures, claims, sections) from the patent HTML.
Usage: registry.py <review.html> <out.xlsx>"""
import re, sys, html as H
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

src, out = sys.argv[1:3]
raw = open(src, encoding="utf-8").read()
soup = BeautifulSoup(raw, "lxml")
wrap = soup.find(class_="wrap")

def added(t):
    while t is not None and hasattr(t, "get"):
        c = t.get("class", []) or []
        if "add" in c or "addfig" in c: return True
        t = t.parent
    return False
def txt(t): return re.sub(r"\s+", " ", t.get_text(" ")).strip()

# ---- figures (raw regex keeps SVG text intact)
figs = []
for m in re.finditer(r'(<div class="addfig"[^>]*>|<div style="page-break-inside:avoid;margin:18px 0 26px">)<div[^>]*>FIG\. (\d+)</div>(<svg.*?</svg>)<div[^>]*>(.*?)</div></div>', raw, re.S):
    num = int(m.group(2)); svg = m.group(3)
    cap = H.unescape(re.sub(r"<[^>]+>", "", m.group(4))).strip()
    title = re.sub(r"^FIG\. \d+\s*[—-]\s*", "", cap)
    texts = [H.unescape(t) for t in re.findall(r"<text[^>]*>(.*?)</text>", svg, re.S)]
    nums = set()
    for t in texts:
        t = re.sub(r"<[^>]+>", "", t)
        nums.update(int(x) for x in re.findall(r"(?<![\d.])(\d{3,4})(?![\d.%])", t))
    figs.append({"n": num, "title": title, "nums": nums, "added": m.group(1).startswith('<div class="addfig"')})
figs.sort(key=lambda f: f["n"])
assert len(figs) == 55, len(figs)

# ---- numerals
tables = wrap.find_all("table")
numtab = [t for t in tables if t.find("th") and "Element" in t.get_text()][0]
numerals = []
for tr in numtab.find_all("tr"):
    tds = tr.find_all("td")
    for i in range(0, len(tds) - 1, 2):
        n, e = txt(tds[i]), txt(tds[i + 1])
        if n.isdigit():
            numerals.append({"n": int(n), "el": e, "added": added(tr)})
numerals.sort(key=lambda x: x["n"])
for x in numerals:
    x["figs"] = [f["n"] for f in figs if x["n"] in f["nums"]]

# ---- sections & paragraphs
sections, cur = [], None
for el in wrap.find_all(["h2", "h3", "p"]):
    if el.name == "h3" and txt(el).startswith("Section "):
        cur = {"title": txt(el), "paras": [], "added": added(el)}; sections.append(cur)
    elif el.name == "h2":
        cur = None
    elif el.name == "p" and cur is not None:
        n = el.find("span", class_="n")
        if n: cur["paras"].append(txt(n))
for s in sections:
    s["figs"] = ", ".join(sorted(set(re.findall(r"FIG\. \d+(?:–FIG\. \d+)?", s["title"]))))

# ---- claims
claims = []
for p in wrap.find_all("p", class_="claim"):
    cn = p.find("span", class_="cn")
    if not cn: continue
    n = int(txt(cn).rstrip("."))
    body = txt(p)[len(txt(cn)):].strip()
    dep = re.match(r"(?:The|A) (?:method|system|subsystem|pipeline|apparatus|medium)[^,]*? of claim (\d+)", body)
    claims.append({"n": n, "dep": int(dep.group(1)) if dep else None, "text": body, "added": added(p)})
claims.sort(key=lambda c: c["n"])

# ================================================================ workbook
F = "Arial"
HDR = PatternFill("solid", start_color="1F3864"); HF = Font(name=F, bold=True, color="FFFFFF", size=10)
YEL = PatternFill("solid", start_color="FFFF00")
BF = Font(name=F, size=10); thin = Side(style="thin", color="BFBFBF")
BOR = Border(left=thin, right=thin, top=thin, bottom=thin)
WR = Alignment(wrap_text=True, vertical="top")

def sheet(ws, headers, widths, rows, added_col):
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.border = HDR, HF, BOR
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        hl = row[added_col].value == "Yes"
        for c in row:
            c.font, c.border, c.alignment = BF, BOR, WR
            if hl: c.fill = YEL
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

wb = Workbook()
ws = wb.active; ws.title = "Summary"

wn = wb.create_sheet("Reference Numerals")
sheet(wn, ["No.", "Element", "Shown in FIG.", "Series", "Added rev. 2026-10-08"], [8, 52, 18, 10, 13],
      [[x["n"], x["el"], ", ".join(map(str, x["figs"])) or "—", f'{x["n"] // 100 * 100}s', "Yes" if x["added"] else "No"] for x in numerals], 4)

wf = wb.create_sheet("Figures")
sheet(wf, ["FIG.", "Title", "Numerals shown", "Numeral count", "Added rev. 2026-10-08"], [7, 70, 60, 10, 13],
      [[f["n"], f["title"], ", ".join(map(str, sorted(x for x in f["nums"] if x in {y["n"] for y in numerals}))),
        None, "Yes" if f["added"] else "No"] for f in figs], 4)
for r in range(2, len(figs) + 2):  # count of listed numerals, computed in-sheet
    wf[f"D{r}"] = f'=IF(LEN(TRIM(C{r}))=0,0,LEN(C{r})-LEN(SUBSTITUTE(C{r},",",""))+1)'

wc = wb.create_sheet("Claims")
sheet(wc, ["Claim", "Type", "Depends on", "Claim text", "Added rev. 2026-10-08"], [7, 12, 10, 110, 13],
      [[c["n"], "Dependent" if c["dep"] else "Independent", c["dep"] if c["dep"] else "—", c["text"], "Yes" if c["added"] else "No"] for c in claims], 4)

wsct = wb.create_sheet("Sections")
sheet(wsct, ["Section", "Paragraphs", "Figures", "Added rev. 2026-10-08"], [90, 22, 18, 13],
      [[s["title"], (f'{s["paras"][0]}–{s["paras"][-1]}' if s["paras"] else "—"), s["figs"] or "—", "Yes" if s["added"] else "No"] for s in sections], 3)

# Summary with live formulas
ws["A1"] = "QueryBook Provisional Patent Application — Registry"; ws["A1"].font = Font(name=F, bold=True, size=14)
ws["A2"] = "Draft rev. October 8, 2026 — not filed. Rows highlighted yellow on the other tabs are additions in this revision."
ws["A2"].font = Font(name=F, italic=True, size=10, color="595959")
ws.append([])
ws.append(["Item", "Total", "Added in rev. 2026-10-08", "Prior draft"])
for c in ws[4]: c.fill, c.font, c.border = HDR, HF, BOR
items = [("Reference numerals", "'Reference Numerals'!A:A", "'Reference Numerals'!E:E"),
         ("Figures", "Figures!A:A", "Figures!E:E"),
         ("Claims", "Claims!A:A", "Claims!E:E"),
         ("Detailed-description sections", "Sections!A:A", "Sections!D:D")]
for i, (name, rng, add) in enumerate(items, 5):
    ws[f"A{i}"] = name
    ws[f"B{i}"] = f"=COUNTA({rng})-1"
    ws[f"C{i}"] = f'=COUNTIF({add},"Yes")'
    ws[f"D{i}"] = f"=B{i}-C{i}"
r = 5 + len(items)
ws[f"A{r}"] = "Independent claims"; ws[f"B{r}"] = '=COUNTIF(Claims!B:B,"Independent")'
ws[f"C{r}"] = '=COUNTIFS(Claims!B:B,"Independent",Claims!E:E,"Yes")'; ws[f"D{r}"] = f"=B{r}-C{r}"
r += 1
ws[f"A{r}"] = "Dependent claims"; ws[f"B{r}"] = '=COUNTIF(Claims!B:B,"Dependent")'
ws[f"C{r}"] = '=COUNTIFS(Claims!B:B,"Dependent",Claims!E:E,"Yes")'; ws[f"D{r}"] = f"=B{r}-C{r}"
for row in ws.iter_rows(min_row=5, max_row=r):
    for c in row: c.font, c.border = BF, BOR
ws[f"A{r+2}"] = ("Numeral-to-figure mapping is read from the drawing labels; “—” means the numeral appears in the "
                 "text and numeral table but not as a label in any drawing.")
ws[f"A{r+2}"].font = Font(name=F, italic=True, size=9, color="595959")
ws.column_dimensions["A"].width = 34
for col in "BCD": ws.column_dimensions[col].width = 22
wb.save(out)
print("numerals", len(numerals), "figs", len(figs), "claims", len(claims), "sections", len(sections),
      "unmapped", sum(1 for x in numerals if not x["figs"]))
