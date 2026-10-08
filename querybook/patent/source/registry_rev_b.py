#!/usr/bin/env python3
"""Build the Excel registry (numerals, figures, claims, sections) from the patent HTML.
Usage: registry.py <review.html> <out.xlsx>"""
import re, sys, html as H
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

src, out = sys.argv[1:3]
raw = open(src, encoding="utf-8").read()
soup = BeautifulSoup(raw, "lxml")
wrap = soup.find(class_="wrap")

def added(t):
    while t is not None and hasattr(t, "get"):
        c = t.get("class", []) or []
        if "add" in c or "addfig" in c: return True
        t = t.parent
    return False
def txt(t): return re.sub(r"\s+", " ", t.get_text(" ")).strip()

# ---- figures (raw regex keeps SVG text intact)
figs = []
for m in re.finditer(r'(<div class="addfig"[^>]*>|<div style="page-break-inside:avoid;margin:18px 0 26px">)<div[^>]*>FIG\. (\d+)</div>(<svg.*?</svg>)<div[^>]*>(.*?)</div></div>', raw, re.S):
    num = int(m.group(2)); svg = m.group(3)
    cap = H.unescape(re.sub(r"<[^>]+>", "", m.group(4))).strip()
    title = re.sub(r"^FIG\. \d+\s*[—-]\s*", "", cap)
    texts = [H.unescape(t) for t in re.findall(r"<text[^>]*>(.*?)</text>", svg, re.S)]
    nums = set()
    for t in texts:
        t = re.sub(r"<[^>]+>", "", t)
        nums.update(re.findall(r"(?<![\w.#-])(\d{3,4}[a-d]?)(?![\w.%])", t))
    figs.append({"n": num, "title": title, "nums": nums, "added": m.group(1).startswith('<div class="addfig"')})
figs.sort(key=lambda f: f["n"])
assert len(figs) == 55, len(figs)

# ---- numerals
tables = wrap.find_all("table")
numtab = [t for t in tables if t.find("th") and "Element" in t.get_text()][0]
numerals = []
for tr in numtab.find_all("tr"):
    tds = tr.find_all("td")
    for i in range(0, len(tds) - 1, 2):
        n, e = txt(tds[i]), txt(tds[i + 1])
        if re.fullmatch(r"\d+[a-d]?", n):
            numerals.append({"n": n, "el": e, "added": tds[i].find(class_="add") is not None})
numerals.sort(key=lambda x: (int(re.match(r"\d+", x["n"]).group(0)), x["n"]))
for x in numerals:
    x["figs"] = [f["n"] for f in figs if x["n"] in f["nums"]]

# ---- sections & paragraphs
sections, cur = [], None
for el in wrap.find_all(["h2", "h3", "h4", "p"]):
    if el.name in ("h3", "h4") and txt(el).startswith("Section "):
        cur = {"title": txt(el), "paras": [], "added": added(el)}; sections.append(cur)
    elif el.name == "h2" or (el.name == "h3" and txt(el).startswith("Part ")):
        cur = None if el.name == "h2" else cur
    elif el.name == "p" and cur is not None:
        n = el.find("span", class_="n")
        if n: cur["paras"].append(txt(n))
for s in sections:
    s["figs"] = ", ".join(sorted(set(re.findall(r"FIG\. \d+(?:–FIG\. \d+)?", s["title"])), key=lambda v: int(re.search(r"\d+", v).group(0))))

# ---- claims
claims = []
for p in wrap.find_all("p", class_="claim"):
    cn = p.find("span", class_="cn")
    if not cn: continue
    n = int(txt(cn).rstrip("."))
    body = txt(p)[len(txt(cn)):].strip()
    dep = re.match(r"(?:The|A) (?:method|system|subsystem|pipeline|apparatus|medium)[^,]*? of claim (\d+)", body)
    claims.append({"n": n, "dep": int(dep.group(1)) if dep else None, "text": body, "added": added(p)})
claims.sort(key=lambda c: c["n"])

# ================================================================ parts
def fig_part(n):
    return "A — Core platform" if n <= 22 else "B — Language & speech" if n <= 32 else "C — Identity, audit, interpretation, lecture" if n <= 40 else "D — Further embodiments & agent-era"
def claim_part(n):
    return "A — Core platform" if n <= 20 else "B — Language & speech" if n <= 50 else "C — Identity, audit, interpretation, lecture" if n <= 82 else "D — Further embodiments & agent-era"
ROMAN = ["I","II","III","IV","V","VI","VII","VIII","IX","X","XI","XII","XIII","XIV","XV","XVI","XVII","XVIII","XIX","XX","XXI","XXII","XXIII"]
def sec_part(title):
    r = re.match(r"Section ([IVX]+)", title).group(1); i = ROMAN.index(r) + 1
    return "A — Core platform" if i <= 8 else "B — Language & speech" if i <= 16 else "C — Identity, audit, interpretation, lecture" if i <= 20 else "D — Further embodiments & agent-era"
num_fig = {}
for x in numerals:
    num_fig[x["n"]] = x["figs"]
def num_part(x):
    if x["figs"]: return fig_part(x["figs"][0])
    return "B — Language & speech"  # 990–993: translator/detector elements of the language subsystem (no drawing label)

# ================================================================ workbook
F = "Arial"
HDR = PatternFill("solid", start_color="1F3864"); HF = Font(name=F, bold=True, color="FFFFFF", size=10)
YEL = PatternFill("solid", start_color="FFFF00")
BF = Font(name=F, size=10); thin = Side(style="thin", color="BFBFBF")
BOR = Border(left=thin, right=thin, top=thin, bottom=thin)
WR = Alignment(wrap_text=True, vertical="top")
NEW = "New in rev. 2026-10-08 (B)"

def sheet(ws, headers, widths, rows, added_col):
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.border = HDR, HF, BOR
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        hl = row[added_col].value == "Yes"
        for c in row:
            c.font, c.border, c.alignment = BF, BOR, WR
            if hl: c.fill = YEL
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

wb = Workbook()
ws = wb.active; ws.title = "Summary"

wn = wb.create_sheet("Reference Numerals")
sheet(wn, ["No.", "Element", "Shown in FIG.", "Part", NEW], [8, 52, 16, 34, 14],
      [[x["n"], x["el"], ", ".join(map(str, x["figs"])) or "—", num_part(x), "Yes" if x["added"] else "No"] for x in numerals], 4)
wf = wb.create_sheet("Figures")
numset = {x["n"] for x in numerals}
sheet(wf, ["FIG.", "Title", "Part", "Numerals shown", "Numeral count", NEW], [7, 66, 34, 56, 10, 14],
      [[f["n"], f["title"], fig_part(f["n"]), ", ".join(sorted((x for x in f["nums"] if x in numset), key=lambda v: (int(re.match(r"\d+", v).group(0)), v))),
        None, "Yes" if f["added"] else "No"] for f in figs], 5)
for r in range(2, len(figs) + 2):
    wf[f"E{r}"] = f'=IF(LEN(TRIM(D{r}))=0,0,LEN(D{r})-LEN(SUBSTITUTE(D{r},",",""))+1)'
wc = wb.create_sheet("Claims")
sheet(wc, ["Claim", "Part", "Type", "Depends on", "Claim text", NEW], [7, 30, 12, 10, 100, 14],
      [[c["n"], claim_part(c["n"]), "Dependent" if c["dep"] else "Independent", c["dep"] if c["dep"] else "—", c["text"], "Yes" if c["added"] else "No"] for c in claims], 5)
wsct = wb.create_sheet("Sections")
sheet(wsct, ["Section", "Part", "Paragraphs", "Figures", NEW], [84, 30, 18, 26, 14],
      [[s["title"], sec_part(s["title"]), (f'{s["paras"][0]}–{s["paras"][-1]}' if s["paras"] else "—"), s["figs"] or "—", "Yes" if s["added"] else "No"] for s in sections], 4)

ws["A1"] = "QueryBook Provisional Patent Application — Registry"; ws["A1"].font = Font(name=F, bold=True, size=14)
ws["A2"] = ("Draft rev. October 8, 2026 (B), rebalanced QueryBook-first — not filed. Rows highlighted yellow on the other tabs are new in this revision.")
ws["A2"].font = Font(name=F, italic=True, size=10, color="595959")
ws.append([])
ws.append(["Balance by part", "Figures", "Claims", "Independent claims", "Sections", "Reference numerals"])
for c in ws[4]: c.fill, c.font, c.border = HDR, HF, BOR
PARTS = ["A — Core platform", "B — Language & speech", "C — Identity, audit, interpretation, lecture", "D — Further embodiments & agent-era"]
for i, p in enumerate(PARTS, 5):
    ws[f"A{i}"] = p
    ws[f"B{i}"] = f'=COUNTIF(Figures!C:C,A{i})'
    ws[f"C{i}"] = f'=COUNTIF(Claims!B:B,A{i})'
    ws[f"D{i}"] = f'=COUNTIFS(Claims!B:B,A{i},Claims!C:C,"Independent")'
    ws[f"E{i}"] = f'=COUNTIF(Sections!B:B,A{i})'
    ws[f"F{i}"] = f"=COUNTIF('Reference Numerals'!D:D,A{i})"
t = 5 + len(PARTS)
ws[f"A{t}"] = "Total"
for col in "BCDEF":
    ws[f"{col}{t}"] = f"=SUM({col}5:{col}{t-1})"
ws[f"A{t}"].font = Font(name=F, bold=True, size=10)
r = t + 2
ws[f"A{r}"] = "New in this revision"; ws[f"A{r}"].font = Font(name=F, bold=True, size=10)
ws[f"B{r}"] = '=COUNTIF(Figures!F:F,"Yes")'; ws[f"C{r}"] = '=COUNTIF(Claims!F:F,"Yes")'
ws[f"D{r}"] = '=COUNTIFS(Claims!C:C,"Independent",Claims!F:F,"Yes")'; ws[f"E{r}"] = '=COUNTIF(Sections!E:E,"Yes")'
ws[f"F{r}"] = "=COUNTIF('Reference Numerals'!E:E,\"Yes\")"
for row in ws.iter_rows(min_row=5, max_row=r):
    for c in row:
        if c.value is not None or c.column > 1: c.border = BOR
        if c.font is None or not c.font.bold: c.font = BF
ws[f"A{t}"].font = Font(name=F, bold=True, size=10); ws[f"A{r}"].font = Font(name=F, bold=True, size=10)
ws[f"A{r+2}"] = ("Numeral-to-figure mapping is read from the drawing labels. Numerals 100 (the LEL system as a whole) and 990–993 (translator and language-detector "
                 "elements) appear in the text and key but are not labelled on any drawing; they are counted under Part B.")
ws[f"A{r+2}"].font = Font(name=F, italic=True, size=9, color="595959")
ws.column_dimensions["A"].width = 44
for col in "BCDEF": ws.column_dimensions[col].width = 16
wb.save(out)
print("numerals", len(numerals), "figs", len(figs), "claims", len(claims), "sections", len(sections),
      "unmapped", [x["n"] for x in numerals if not x["figs"]])
