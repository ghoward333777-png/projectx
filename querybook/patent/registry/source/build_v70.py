#!/usr/bin/env python3
"""Composite Feature Registry v70 = v69 (every row and sheet unchanged) + the Self-Knowledge Layer:
features 383-391 (FIG. 57, numerals 5710-5790), algorithms ALG-20 and ALG-21, drawing FIG. 57, updated summary and record.
Usage: build_v70.py <v69.xlsx> <out.xlsx>"""
import copy, os, re, sys
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "docs", "source"))
import skl_text as K

src, out = sys.argv[1:3]
wb = load_workbook(src)
ref = load_workbook(src)                      # untouched copy for the additive check
ws = wb["Feature Registry v69"]; ws.title = "Feature Registry v70"
wa = wb["Algorithms (v69)"]; wa.title = "Algorithms (v70)"
wd = wb["Drawings FIG 1-56"]; wd.title = "Drawings FIG 1-57"
sm = wb["Summary v69"]; sm.title = "Summary v70"
REG = "'Feature Registry v70'"
NEWFILL = PatternFill("solid", start_color="FFFF00")

def style_like(src_cell, c, fill=None):
    c.font, c.border = copy.copy(src_cell.font), copy.copy(src_cell.border)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if fill: c.fill = fill

# ------------------------------------------------ new feature rows
OLD_LAST = ws.max_row
tpl = ws.cell(OLD_LAST, 2)
first_no = int(ws.cell(OLD_LAST, 1).value) + 1
for k, ((name, dom, figs, surface, role, func, rule, wexpr, status, module), n) in enumerate(zip(K.NEW, K.NUMERAL)):
    r = OLD_LAST + 1 + k; no = first_no + k
    cite = f"{K.CITE} · FIG. " + ", ".join(map(str, figs))
    func = f"{func[:-1]} ({n})."
    vals = [no, name, dom, cite, f"{surface}{cite}{role} {func}", rule, wexpr, ", ".join(map(str, figs)), status, module, "v70"]
    for col, v in enumerate(vals, 1):
        style_like(tpl, ws.cell(r, col, v), NEWFILL)
LAST = ws.max_row
ws.auto_filter.ref = f"A1:K{LAST}"

# ------------------------------------------------ algorithms
atpl = wa.cell(wa.max_row, 2)
for aid, name, feats, figs, inp, proc, outp, det, mod in K.ALGOS:
    r = wa.max_row + 1
    for i, v in enumerate([aid, name, feats, ", ".join(map(str, figs)), inp, proc, outp, det, mod], 1):
        style_like(atpl, wa.cell(r, i, v))
wa.auto_filter.ref = f"A1:I{wa.max_row}"
ALAST = wa.max_row

# ------------------------------------------------ drawings: FIG. 57, and every formula re-pointed at the renamed sheets
dtpl = wd.cell(wd.max_row, 2)
r = wd.max_row + 1
for i, v in enumerate([K.FIG, K.FIG_TITLE, "Self-knowledge"], 1):
    style_like(dtpl, wd.cell(r, i, v))
for rr in range(2, r + 1):
    wd.cell(rr, 4).value = f'=SUMPRODUCT(--ISNUMBER(SEARCH(", "&A{rr}&",",", "&{REG}!$H$2:$H${LAST}&",")))'
    wd.cell(rr, 5).value = f"=SUMPRODUCT(--ISNUMBER(SEARCH(\", \"&A{rr}&\",\",\", \"&'Algorithms (v70)'!$D$2:$D${ALAST}&\",\")))"
    if rr == r:
        for c in (4, 5): style_like(dtpl, wd.cell(rr, c))

# ------------------------------------------------ summary
BF = sm["A5"].font
sm["A1"] = "QueryBook Composite Feature Registry v70"
sm["A2"] = ("Rows highlighted yellow on the Feature Registry sheet were added in v66, v69 or v70 (column K). Verified Agent Transport is "
            "supported, not included: see the Supported External Services sheet. All v65 rows (columns A–G) and all v66 and v69 rows are unchanged.")
for row in sm.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith("="):
            c.value = (c.value.replace("'Feature Registry v69'", REG).replace(f"2:A{OLD_LAST}", f"2:A{LAST}")
                       .replace(f"2:C{OLD_LAST}", f"2:C{LAST}").replace(f"2:G{OLD_LAST}", f"2:G{LAST}").replace(f"2:H{OLD_LAST}", f"2:H{LAST}")
                       .replace(f"2:I{OLD_LAST}", f"2:I{LAST}").replace(f"2:K{OLD_LAST}", f"2:K{LAST}")
                       .replace("'Algorithms (v69)'", "'Algorithms (v70)'").replace("A2:A20", f"A2:A{ALAST}")
                       .replace("'Drawings FIG 1-56'", "'Drawings FIG 1-57'").replace("A2:A57", "A2:A58").replace("D2:D57", "D2:D58"))
        elif c.value == "Drawings (FIG. 1–56)":
            c.value = "Drawings (FIG. 1–57)"
# a row for v70 after the v69 row, by shifting the rows below it
vrow = next(r for r in range(1, sm.max_row + 1) if str(sm.cell(r, 1).value).startswith("Features added in v69"))
sm.insert_rows(vrow + 1)
sm.cell(vrow + 1, 1, "Features added in v70 (Self-Knowledge Layer, FIG. 57)").font = copy.copy(BF)
sm.cell(vrow + 1, 2, f'=COUNTIF({REG}!K2:K{LAST},"v70")').font = copy.copy(BF)
hdr = next(r for r in range(1, sm.max_row + 1) if sm.cell(r, 1).value == "Features by domain")
sm.cell(hdr, 5, "New in v70").font = copy.copy(sm.cell(hdr, 4).font)
r = hdr + 1
while sm.cell(r, 1).value:
    for c in (2, 3, 4):                                   # insert_rows does not move formula references: re-point A{r}
        v = sm.cell(r, c).value
        if isinstance(v, str):
            sm.cell(r, c).value = re.sub(r"\bA\d+\b", f"A{r}", v)
    sm.cell(r, 5, f'=COUNTIFS({REG}!C2:C{LAST},A{r},{REG}!K2:K{LAST},"v70")').font = copy.copy(BF)
    r += 1
sm.column_dimensions["E"].width = 12

# ------------------------------------------------ registry record
rec = wb["Registry Record"]
for l in ["",
          "QueryBook Composite Feature Registry v70",
          f"Assembled {K.DATE} from Composite Registry v69 and synchronized with the provisional application as updated with the "
          f"Self-Knowledge Layer (FIG. 57; claims 358–362) and with Core Technology Overview v15 §56.",
          f"Features: {LAST - 1}. Rows 1–{OLD_LAST - 1} carried from v69 unchanged. Rows {first_no}–{LAST - 1} added in v70: the "
          f"Self-Knowledge Layer, each citing FIG. 57 and its reference numeral (5710–5790).",
          f"Sheets: Feature Registry v70, Algorithms (v70) — {ALAST - 1} algorithms (ALG-20 and ALG-21 are self-knowledge algorithms), "
          "Drawings FIG 1-57, Supported External Services, Summary v70.",
          "Status of all v70 rows: Built — prototype v9.64 (qb_selfknow.py) and the applications' shared core (qbselfknow.py). "
          "Additive covenant: no feature removed; no v65 cell in columns A–G altered; no v66 or v69 row altered. Claims: 362."]:
    rec.append([l])
for row in rec.iter_rows():
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")

# ------------------------------------------------ additive check: every v69 feature cell unchanged
a = ref["Feature Registry v69"]
diff = [(r, c) for r in range(1, OLD_LAST + 1) for c in range(1, 12) if a.cell(r, c).value != ws.cell(r, c).value]
assert not diff, diff[:5]
wb.active = 0
wb.save(out)
print("features:", LAST - 1, "new:", LAST - OLD_LAST, f"({first_no}-{LAST - 1})", "algorithms:", ALAST - 1)
