import zipfile, shutil, re, html, sys
from openpyxl import load_workbook
a=load_workbook(sys.argv[1])['Feature Registry v65']; src=sys.argv[2]; sheet='xl/worksheets/sheet2.xml'
z=zipfile.ZipFile(src); x=z.read(sheet).decode(); fixed=[]
for r in range(2,a.max_row+1):
    v=a.cell(r,7).value
    if isinstance(v,str) and v.startswith('='):
        m=re.search(r'<c r="G%d"[^>]*><f[^>]*>(.*?)</f>'%r,x); assert m, r
        if html.unescape(m.group(1))!=v[1:]:
            x=x[:m.start(1)]+html.escape(v[1:],quote=False)+x[m.end(1):]; fixed.append(r)
with zipfile.ZipFile(src+'.tmp','w',zipfile.ZIP_DEFLATED) as o:
    for n in z.namelist(): o.writestr(z.getinfo(n), x if n==sheet else z.read(n))
z.close(); shutil.move(src+'.tmp',src)
b=load_workbook(src)['Feature Registry v66']
diff=[(r,c) for r in range(1,a.max_row+1) for c in range(1,8) if a.cell(r,c).value!=b.cell(r,c).value]
print('restored rows:',fixed,'| v65 cells A–G changed:',diff)
