"""Add FIG. 57 (Self-Knowledge Layer) to the Drawings and Description document: brief description, detailed description,
reference numerals and status, in the styles of the FIG. 56 entries. Usage: desc_skl.py <FIG1-56.docx> <FIG1-57.docx>"""
import copy, os, sys
from docx import Document
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "..", "docs", "source"))
import skl_text as T

src, out = sys.argv[1:3]
d = Document(src); P = d.paragraphs; body = d.element.body
find = lambda s: next(p for p in P if p.text.startswith(s))
h2, brief, h1 = find("Support for agent-wrapper services (FIG. 56)"), find("FIG. 56 is a block diagram"), find("Detailed Description of FIG. 56")
pbr = h1._p.getprevious(); txt = find("FIG. 56 shows the interfaces"); stat = find("The support interfaces of FIG. 56 are specified")
tbl = d.tables[0]

def para(tpl, runs):
    p = copy.deepcopy(tpl._p)
    rs = p.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r")
    for r in rs[1:]: p.remove(r)
    r0 = rs[0]
    from docx.text.paragraph import Paragraph
    para_ = Paragraph(p, tpl._parent)
    para_.runs[0].text = runs[0][0]
    for s, bold in runs[1:]:
        r = copy.deepcopy(r0); p.append(r); Paragraph(p, tpl._parent).runs[-1].text = s
        Paragraph(p, tpl._parent).runs[-1].bold = bold
    para_.runs[0].bold = runs[0][1]
    return p

# brief description after FIG. 56
a = brief._p
for el in [para(h2, [(f"Self-knowledge (FIG. {T.FIG})", None)]), para(brief, [(t, b) for t, b in T.BRIEF[0]])]:
    a.addnext(el); a = el
# detailed description at the end
sect = body[-1]
els = [copy.deepcopy(pbr), para(h1, [(f"Detailed Description of FIG. {T.FIG}", None)])]
els += [para(txt, [(t, None)]) for t in [T.INTRO[0], T.CORE[0][1][0]]]
els.append(para(find("Reference numerals"), [("Reference numerals", None)]))
nt = copy.deepcopy(tbl._tbl)
rows = nt.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr")
for r in rows[2:]: nt.remove(r)
els.append(nt)
els += [para(find("Status"), [("Status", None)]),
        para(stat, [(f"The self-knowledge layer of FIG. {T.FIG} is built and tested in prototype v9.64 (qb_selfknow.py) and in the shared core "
                     "of the applications (qbselfknow.py). It corresponds to features 383–391 of the QueryBook Composite Feature Registry v70, "
                     "to Section 56 of the Core Technology Overview v15 and to claims 358–362 of the application.", None)])]
for el in els: sect.addprevious(el)
# fill the numeral table: header row kept, first data row as the template
from docx.table import Table
t = Table(nt, d._body)
tpl_row = t.rows[1]._tr
for n, (num, what) in enumerate(T.FIG_KEY):
    tr = tpl_row if n == 0 else copy.deepcopy(tpl_row)
    if n: nt.append(tr)
    from docx.table import _Row
    row = _Row(tr, t)
    rpr = row.cells[0].paragraphs[0].runs[0]._r.rPr
    for c, v in zip(row.cells, [num, what, "FIG. 57"]):
        q = c.paragraphs[0]
        if not q.runs:
            q.add_run("")
            if rpr is not None: q.runs[0]._r.insert(0, copy.deepcopy(rpr))
        q.runs[0].text = v
        for r in q.runs[1:]: r.text = ""
# titles
P[1].runs[0].text = "FIG. 1–57 · companion to the Provisional Patent Application (synchronized, with support for agent-wrapper services and the self-knowledge layer)"
for r in P[1].runs[1:]: r.text = ""
P[2].runs[0].text = T.DATE
for s in d.sections:
    for p in s.header.paragraphs:
        for r in p.runs: r.text = r.text.replace("FIG. 1–56", "FIG. 1–57")
d.save(out)
print("saved", out)
