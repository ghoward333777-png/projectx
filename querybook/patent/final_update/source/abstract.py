from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_COLOR_INDEX, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import sys
HL = sys.argv[2] == "1"
d = Document()
s = d.sections[0]; s.page_width, s.page_height = Inches(8.5), Inches(11)
for m in ("left_margin","right_margin","top_margin","bottom_margin"): setattr(s, m, Inches(1))
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(10)
h = d.add_paragraph(); r = h.add_run("Abstract"); r.bold = True
P1 = "QueryBook represents a new model of AI - a synthetic cortex and deterministic world‑model AI engine built on a new primitive—the Fact Unit, an atomic sub-language semantic unit. QueryBook’s engine provides neuromorphic reasoning, multimodal cognition, and an open, flexible architecture. It represents a new paradigm in AI—one that is structured, explainable, multimodal, and universally deployable across virtually every current AI application. QueryBook’s ability can be micro-encapsulated for Q&A of single or multiple books and documents in smart devices, used for data analysis, and provide SaaS, ChatBot, Agent, deterministic hybrid additions to current LLMs, specialized  LLM-like applications including private AI computational and analysis systems (within option for live data inputs) for those with specialized knowledge, plus unique Fact Unit compressed data search engine capability that can operate alone for specialized fields or as adjunct to current search engines to off load centralized data computation, storage and transmission costs -- bringing to each application a distinct and measurable set of advantages. "
P2 = " These deployment models have modular capability editions (light, medium, and full) selectable at launch, and with a version-independent, provenance-continuous data store that keeps collected Fact Units in one fixed location so that software upgrades never lose or orphan previously harvested facts. QueryBook transforms digital content into a coherent, navigable, intelligent world‑model, enabling users to understand, explore, and apply knowledge with unprecedented clarity and precision with option for speech recognition and natural language response that includes emotional inflection as if hearing a human voice-- deterministic, non-hallucinogenic, with citations. Across all applications, internal controls ensure humanistic action and prohibition of rogue actions, and its innovative multi- layer security infrastructure provides a model for AI and computer systems offering an anti-intrusion shield against viral attacks including Agent Worms. Further additions strengthen this security and integrity posture: an independent Universal Integrity Audit Service (UIAS) produces pre-dispute, tamper-evident proof over the content-addressed Fact Unit ledger for a third party's system, returning only a factual determination (verified, contradicted, or unknown) and never a legal conclusion; and a key-safe external-service mediation layer holds any external-service credential on the server side of the query/presentation boundary so it is never exposed to a client."
A1 = " Recent additions extend this protection to AI agents: every agent call passes a single gateway with traps and an AI-aware firewall, a hostile caller is isolated in a decoy with no path to real data, security decisions are sealed in a tamper-evident log, and Fact Units may optionally be certified on a public blockchain for independent verification."
P3a = "QueryBook Fact Unit graph structure simultaneously permits a (a) new level of enhanced multi-language provision and translation that includes deterministic, dictionary-grounded multilingual acquisition and offline speech with no large language model in the grounding path, regional Dialect Parameter Clusters, a Lecture Query mode that ingests a live or recorded talk as attributed utterances (never promoted to world-truth), (b) Scene Director and Scene Reconstructor that assemble grounded video prompts and reconstruct scenes from prose with all generated media clearly marked as reconstruction, (c) simplified colorization capability, (d) video and audio compression benefits for security applications, video expansion (not yet tested) and, (e) a new Fact Unit based, enhanced TCP/IP hybrid protocol carrying Fact Unit identifiers and provenance references alongside network payload data."
A2 = " Files of any format can be ingested as data, never as instructions, and external AI agents such as OpenClaw connect through a standard tool interface limited to their allow-list. "
P3b = "QueryBooks Fact Unit graph architecture also reduces data storage by orders of magnitude, works on commodity server or product hardware, and in larger complex application significantly reduce power consumption, cooling costs and deployment real estate."
def para(parts):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for t, new in parts:
        r = p.add_run(t)
        if new and HL: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
para([(P1, False)]); para([(P2, False), (A1, True)]); para([(P3a, False), (A2, True), (P3b, False)])
d.core_properties.title = "QueryBook Abstract (updated)"
z = d.settings.element.find(qn("w:zoom"))
if z is not None and z.get(qn("w:percent")) is None: z.set(qn("w:percent"), "100")
d.save(sys.argv[1])
